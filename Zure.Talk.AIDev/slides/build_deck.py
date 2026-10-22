"""Build the talk deck on the Zure template.

    python slides/build_deck.py

Layouts, wordmark and fonts come from slides/template.pptx (the Zure template
with its slides dropped). Speaker notes come from docs/Presentation.md, so edit
notes there and rebuild. Output: docs/AI-Assisted Development, Live.pptx
"""
import os
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "slides" / "template.pptx"
NOTES_MD = ROOT / "docs" / "Presentation.md"
OUT = Path(os.environ.get("DECK_OUT", ROOT / "docs" / "AI-Assisted Development, Live.pptx"))  # DECK_OUT: build elsewhere while the deck is open
ASSETS = Path(__file__).resolve().parent / "assets"

FONT = "Segoe UI"  # what the Zure template uses throughout
MONO = "Consolas"  # commands and file names
INK = RGBColor(0x15, 0x17, 0x19)
INK2 = RGBColor(0x55, 0x5A, 0x5E)
MUTED = RGBColor(0x8A, 0x90, 0x95)
LINE = RGBColor(0xDD, 0xDA, 0xD2)
PAPER = RGBColor(0xF6, 0xF5, 0xF1)
CARD_DARK = RGBColor(0x25, 0x28, 0x2B)
CYAN = RGBColor(0x05, 0xC3, 0xDE)
CYAN_LIGHT = RGBColor(0x7D, 0xEC, 0xFC)
RED = RGBColor(0xFB, 0x52, 0x3C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY_ON_DARK = RGBColor(0xB9, 0xBE, 0xC2)

W, H = 13.333, 7.5
M = 0.6  # side margin, matches the template's title inset


# --- notes from Presentation.md ------------------------------------------------

def load_notes():
    notes, key, buf, in_notes = {}, None, [], False
    for line in NOTES_MD.read_text(encoding="utf-8").splitlines():
        head = re.match(r"^#{2,3} (\d+|A\d+)\. ", line)
        if head or line.strip() == "---":
            if key and buf:
                notes[key] = "\n".join(buf).strip()
            key, buf, in_notes = (head.group(1) if head else None), [], False
            continue
        if line.startswith("**Notes:**"):
            in_notes = True
            continue
        if in_notes and key:
            buf.append(line)
    if key and buf:
        notes[key] = "\n".join(buf).strip()
    return notes


# --- helpers -------------------------------------------------------------------

prs = Presentation(TEMPLATE)
for sld in list(prs.slides._sldIdLst):
    prs.part.drop_rel(sld.rId)
    prs.slides._sldIdLst.remove(sld)

LAYOUTS = {l.name: l for m in prs.slide_masters for l in m.slide_layouts}

# The template's notes master has no placeholders, so notes slides would have no
# notes body. Give it python-pptx's default notes-master shapes.
from pptx.parts.slide import NotesMasterPart  # noqa: E402

_nm_tree = prs.notes_master.shapes._spTree
if not prs.notes_master.placeholders:
    _default = NotesMasterPart.create_default(prs.part.package).notes_master
    for sp in list(_default.shapes._spTree)[2:]:  # skip nvGrpSpPr and grpSpPr
        _nm_tree.append(sp)
NOTES = load_notes()


def slide(layout, key, keep_title=False):
    s = prs.slides.add_slide(LAYOUTS[layout])
    for ph in list(s.placeholders):
        if keep_title and ph.placeholder_format.idx == 0:
            continue
        ph._element.getparent().remove(ph._element)
    if key in NOTES:
        s.notes_slide.notes_text_frame.text = NOTES[key]
    return s


def text(s, x, y, w, h, runs, size=24, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None, name=None):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, opts)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", FONT)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return tb


def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, weight=1.25, name=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sh.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(weight)
    sh.shadow.inherit = False
    return sh


def line(s, x1, y1, x2, y2, color=MUTED, weight=1.25, dash=False, arrow=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(weight)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    if arrow:
        etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd")).set("type", "triangle")
    return c


def arrow(s, x, y, w=0.3, h=0.26, color=LINE):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    return a


def title(s, t, size=40, color=None):
    ph = s.shapes.title
    ph.text = t
    for p in ph.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
    return ph


def kicker(s, label, dark=False):
    """Small cyan label top-left, the template's section-label style."""
    text(s, M, 0.55, 12, 0.3, label.upper(), size=14, bold=True, color=CYAN if not dark else CYAN_LIGHT)


def icon(s, name, x, y, size):
    s.shapes.add_picture(str(ASSETS / f"{name}.png"), Inches(x), Inches(y), Inches(size), Inches(size))


def statement(s, t, y=2.6, size=60, color=INK, w=W - 2 * M, x=M, align=PP_ALIGN.LEFT):
    return text(s, x, y, w, 2.4, t, size=size, bold=True, color=color, align=align, anchor=MSO_ANCHOR.MIDDLE)


def loop(s, cx, cy, r, lit, dark=False, labels=("Spec", "Test", "Harness", "Skill")):
    """The talk's loop: a ring with four numbered nodes, clockwise from the top."""
    ring = box(s, cx - r, cy - r, 2 * r, 2 * r, line=CYAN if lit == "all" else (MUTED if not dark else INK2),
               shape=MSO_SHAPE.OVAL, weight=3, name="Loop ring")
    pos = [(cx, cy - r), (cx + r, cy), (cx, cy + r), (cx - r, cy)]
    d = 0.62
    for i, ((x, y), lab) in enumerate(zip(pos, labels)):
        on = lit == "all" or lit == i
        node = box(s, x - d / 2, y - d / 2, d, d, fill=CYAN if on else (WHITE if not dark else CARD_DARK),
                   line=CYAN if on else MUTED, shape=MSO_SHAPE.OVAL, weight=2, name=f"Loop node {i + 1}")
        tf = node.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        rr = p.add_run()
        rr.text = str(i + 1)
        rr.font.name, rr.font.size, rr.font.bold = FONT, Pt(20), True
        rr.font.color.rgb = INK if on else (MUTED if not dark else GREY_ON_DARK)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        lw, lh = 2.6, 0.6
        lx, ly, al = {
            0: (x - lw / 2, y - d / 2 - lh - 0.08, PP_ALIGN.CENTER),
            1: (x + d / 2 + 0.2, y - lh / 2, PP_ALIGN.LEFT),
            2: (x - lw / 2, y + d / 2 + 0.08, PP_ALIGN.CENTER),
            3: (x - d / 2 - 0.2 - lw, y - lh / 2, PP_ALIGN.RIGHT),
        }[i]
        text(s, lx, ly, lw, lh, lab, size=28, bold=True, align=al, anchor=MSO_ANCHOR.MIDDLE,
             color=(INK if not dark else WHITE) if on else (INK2 if not dark else GREY_ON_DARK))
    return ring


def section(key, num, name, line_):
    """A segment's header slide: big number, phase name, one line."""
    s = slide("Dark: Basic Blue 2", key)
    text(s, M, 1.5, 8, 1.6, num, size=120, bold=True, color=CYAN_LIGHT)
    text(s, M, 3.1, 10, 1.6, name, size=96, bold=True, color=WHITE)
    text(s, M, 4.85, 10, 0.6, line_, size=28, color=GREY_ON_DARK)
    return s


def scorecard(s, cols, rows, x, y, w, row_h=0.85, first_w=4.0, dim=(), head_size=24, label_size=22, baseline=None):
    """A quiet table: no fills, hairlines between rows, values left as dashes to fill live.
    dim: column names to grey out; baseline: the column name that gets a cyan label."""
    if baseline:
        cw = (w - first_w) / len(cols)
        text(s, x + first_w + cols.index(baseline) * cw, y - 0.35, cw, 0.3, "BASELINE", size=12, bold=True,
             color=CYAN, align=PP_ALIGN.CENTER)
    shape = s.shapes.add_table(len(rows) + 1, len(cols) + 1, Inches(x), Inches(y), Inches(w), Inches(row_h * (len(rows) + 1)))
    shape.name = "Scorecard"
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow"):
        tblPr.set(a, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    tbl.columns[0].width = Inches(first_w)
    for c in range(1, len(cols) + 1):
        tbl.columns[c].width = Emu(int(Inches((w - first_w) / len(cols))))

    def cell(r, c, t, size, bold, color, align):
        ce = tbl.cell(r, c)
        ce.fill.background()
        ce.vertical_anchor = MSO_ANCHOR.MIDDLE
        ce.margin_left = Inches(0.1) if c else 0
        ce.margin_right = Inches(0.1)
        tf = ce.text_frame
        tf.paragraphs[0].alignment = align
        rr = tf.paragraphs[0].add_run()
        rr.text = t
        rr.font.name, rr.font.size, rr.font.bold = FONT, Pt(size), bold
        rr.font.color.rgb = color
        tcPr = ce._tc.get_or_add_tcPr()
        for pos, side in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
            ln = etree.Element(qn(side))
            tcPr.insert(pos, ln)  # borders must precede the fill in tcPr
            if side == "a:lnB":
                ln.set("w", str(Pt(1)))
                etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr")).set("val", "DDDAD2")
            else:
                ln.set("w", "0")
                etree.SubElement(ln, qn("a:noFill"))

    cell(0, 0, "", 14, False, INK, PP_ALIGN.LEFT)
    for c, name in enumerate(cols, 1):
        cell(0, c, name, head_size, True, MUTED if name in dim else INK, PP_ALIGN.CENTER)
    for r, label in enumerate(rows, 1):
        cell(r, 0, label, label_size, False, INK2, PP_ALIGN.LEFT)
        for c, name in enumerate(cols, 1):
            cell(r, c, "—", 28, True, LINE if name in dim else MUTED, PP_ALIGN.CENTER)
    return shape


# --- the deck ------------------------------------------------------------------
# Keys match the slide numbers in docs/Presentation.md, where the notes live.

# 1. Our job has changed: the opener, alone on white
s = slide("Light: Empty", "1")
statement(s, "Our job has changed", size=96, y=2.55)

# 2. Silver bullet: alone on white; screenshots of speakers get added by hand
s = slide("Light: Empty", "2")
statement(s, "I wish I had a silver bullet", size=96, y=2.55)

# 3. Idiom: alone on white, like the two before it
s = slide("Light: Empty", "3")
statement(s, "Peeing your pants in winter", size=96, y=2.55)

# 4. Title
s = slide("Dark: Title and Subtitle", "4", keep_title=True)
title(s, "Engineering with AI", size=54)
# A newscast's LIVE caption above the title
bw = 1.5
box(s, (W - bw) / 2, 1.6, bw, 0.5, fill=RED, radius=0.2, name="Live badge")
text(s, (W - bw) / 2, 1.6, bw, 0.5, "●  LIVE", size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 2.0, 4.3, W - 4.0, 0.5, "(developers are stronger than ever)", size=26, color=WHITE, align=PP_ALIGN.CENTER)
text(s, 2.0, 5.1, W - 4.0, 0.5, "Sakari Nahi · Zure", size=20, color=GREY_ON_DARK, align=PP_ALIGN.CENTER)

# 5. 185k€: one colleague's tokens make the case for managing context
s = slide("Dark: Empty", "5")
text(s, M, 1.3, 11, 3.8, "185k€", size=220, bold=True, color=CYAN, anchor=MSO_ANCHOR.MIDDLE)
text(s, M + 0.15, 5.25, 11, 0.9, [[("one colleague's tokens, January to August", {"size": 28, "color": WHITE})],
                                  [("about 23k€ a month", {"size": 16, "color": GREY_ON_DARK})]])

# 6. The gauntlet loop: a harsh critic as the sensor
s = slide("Light: Empty", "6")
text(s, M, 0.9, 12, 0.8, "Loop until a harsh critic is wowed", size=36, bold=True)
PROMPT = ["Fan out sub-agents …", "/loop on each item and have a separate", "sub-agent check it visually …",
          "That separate sub-agent should be", "a really harsh critic …", "Don't stop until each sub-agent is",
          "utterly wowed … compare them", "side by side blind …"]
box(s, M, 2.1, 5.8, 3.6, fill=CARD_DARK, radius=0.04, name="Prompt")
text(s, M + 0.3, 2.1, 5.2, 3.6, [[(t, {"font": MONO})] for t in PROMPT], size=17, color=GREY_ON_DARK, spacing=4,
     anchor=MSO_ANCHOR.MIDDLE)
# The loop: builders into the critic, and back until it is wowed
for x, label, fill, ink in ((7.0, "builders", PAPER, INK), (10.3, "harsh critic", INK, WHITE)):
    box(s, x, 2.3, 2.4, 1.0, fill=fill, line=None if fill == INK else LINE, name=label)
    text(s, x, 2.3, 2.4, 1.0, label, size=20, bold=True, color=ink, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
arrow(s, 9.6, 2.67, 0.5, 0.26, color=MUTED)
line(s, 11.5, 3.35, 11.5, 3.75, color=MUTED, dash=True)
line(s, 11.5, 3.75, 8.2, 3.75, color=MUTED, dash=True)
line(s, 8.2, 3.75, 8.2, 3.35, color=MUTED, dash=True, arrow=True)
text(s, 7.0, 3.85, 5.7, 0.4, "not AAA yet: keep going", size=14, color=MUTED, align=PP_ALIGN.CENTER)
text(s, 7.0, 4.45, 5.7, 0.9, "3.59 → 5.05", size=48, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
text(s, 7.0, 5.35, 5.7, 0.7, [[("out of 10 from the critics", {})], [("blind A/B: the real Call of Duty won every time", {})]],
     size=14, color=INK2, align=PP_ALIGN.CENTER)
text(s, M, 6.55, 12, 0.4, [[("Matt Shumer · Claude of Duty  ", {"bold": True, "color": INK2}),
                            ("github.com/mshumer/Claude-of-Duty", {"color": MUTED})]], size=16)

# 7. The process: spec, implement, QA, deploy, with harness and skills throughout
s = slide("Light: Empty", "7")
kicker(s, "The process")
text(s, M, 0.9, 12, 0.8, "The process we always had, with new techniques", size=32, bold=True)
PHASES = (("Spec", "Say what to build and why, before any code exists.", ("spec-driven development", "context engineering")),
          ("Implement", "The agent builds in small steps, inside guides and sensors.", ("harness engineering", "context engineering")),
          ("QA", "Prove it works with tests and numbers, not vibes.", ("agentic TDD", "mutation testing", "critic agents")),
          ("Deploy", "Put it in front of real people, and watch them and the bill.", ("harness engineering", "cost and context control")))
pg = 0.3
pw = (W - 2 * M - 3 * pg) / 4
d = 0.5
for i, (name, what, how) in enumerate(PHASES):
    x = M + i * (pw + pg)
    node = box(s, x, 2.0, d, d, fill=CYAN, line=CYAN, shape=MSO_SHAPE.OVAL, weight=2, name=f"Phase {i + 1}")
    text(s, x, 2.0, d, d, str(i + 1), size=18, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + d + 0.15, 2.0, pw - d - 0.15, d, name, size=24, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if i < 3:
        line(s, x + pw - 0.12, 2.25, x + pw + pg - 0.05, 2.25, color=MUTED, arrow=True)
    text(s, x, 2.7, pw, 0.9, what, size=16, color=INK2)
    text(s, x, 3.55, pw, 0.9, list(how), size=14, bold=True, color=CYAN)
line(s, M, 4.5, W - M, 4.5, color=LINE)
TERMS = (("Spec-driven development", "The spec, not the chat, is the source of truth the agent builds from and is checked against."),
         ("Context engineering", "Choosing what goes into the model’s window, and keeping everything else out."),
         ("Harness engineering", "Building the guides (rules, skills) and sensors (tests, linters, hooks) around the model."),
         ("Agentic TDD", "The agent writes a failing test first, then code until it passes. The tests say when it’s done."))
for i, (term, means) in enumerate(TERMS):
    x = M + i * (pw + pg)
    text(s, x, 4.75, pw, 1.6, [[(term, {"bold": True, "color": INK, "size": 16})], [(means, {})]],
         size=14, color=INK2, spacing=6)
text(s, M, 6.55, 12, 0.4, "Harness and skills run through all four: anything you had to say twice becomes a skill.",
     size=14, color=MUTED)

# 8. Section: Spec
section("8", "01", "Spec", "Say what to build and why, before any code exists")

# 9. Spec
s = slide("Dark: Basic Blue 2", "9")
kicker(s, "01 · Spec", dark=True)
statement(s, "Say what you want, not how to build it", size=60, color=WHITE, w=8.0)

# 10. OpenSpec
s = slide("Light: Empty", "10")
kicker(s, "01 · Spec")
text(s, M, 0.9, 12, 0.9, "OpenSpec makes the spec step a handful of commands", size=32, bold=True)
init_w, nw, ng, ny, nh = 1.55, 1.78, 0.32, 3.0, 0.95
box(s, M, ny + 0.1, init_w, nh - 0.2, fill=None, line=MUTED, name="openspec init")
text(s, M, ny + 0.1, init_w, nh - 0.2, [[("openspec init", {"font": MONO})]], size=13, color=INK2,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, M, ny + nh + 0.12, init_w, 0.5, "once per project", size=13, color=MUTED, align=PP_ALIGN.CENTER)
fx0 = M + init_w + ng
line(s, M + init_w + 0.04, ny + nh / 2, fx0 - 0.04, ny + nh / 2, color=MUTED, dash=True, arrow=True)
FLOW = (("/opsx:explore", "understand the problem and the code"),
        ("/opsx:propose", "write the change down"),
        ("/opsx:apply", "build the tasks"),
        ("/opsx:verify", "check the build against the spec"),
        ("/opsx:archive", "file the finished change"))
for i, (cmd, cap) in enumerate(FLOW):
    x = fx0 + i * (nw + ng)
    box(s, x, ny, nw, nh, fill=PAPER, line=CYAN if i == 1 else None, weight=2, name=f"OpenSpec {cmd}")
    text(s, x, ny, nw, nh, [[(cmd, {"font": MONO})]], size=15, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x, ny + nh + 0.12, nw, 0.7, cap, size=14, color=INK2, align=PP_ALIGN.CENTER)
    if i < len(FLOW) - 1:
        arrow(s, x + nw + (ng - 0.22) / 2, ny + nh / 2 - 0.11, 0.22, 0.22, color=CYAN)
# What propose writes: four files fanned out under it
px = fx0 + nw + ng
for j, f in enumerate(("proposal.md", "specs/", "design.md", "tasks.md")):
    fx, fy = px + (j % 2) * (nw / 2 + 0.03) - 0.02, 4.75 + (j // 2) * 0.48
    box(s, fx, fy, nw / 2 - 0.02, 0.4, fill=WHITE, line=LINE, shape=MSO_SHAPE.FOLDED_CORNER, name=f"File {f}")
    text(s, fx, fy, nw / 2 - 0.02, 0.4, [[(f, {"font": MONO})]], size=9, color=INK2, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
# Archive loops back to explore for the next change
ax, ex, top = fx0 + 4 * (nw + ng) + nw / 2, fx0 + nw / 2, ny - 0.45
line(s, ax, ny, ax, top, color=MUTED)
line(s, ax, top, ex, top, color=MUTED)
line(s, ex, top, ex, ny - 0.02, color=MUTED, arrow=True)
text(s, ex, top - 0.4, ax - ex, 0.35, "next change", size=13, color=MUTED, align=PP_ALIGN.CENTER)
text(s, M, 6.55, 12, 0.4, [[("npm install -g @fission-ai/openspec@latest", {"font": MONO})]], size=13, color=MUTED)

# 11. Section: Implement
section("11", "02", "Implement", "The agent builds, inside guides and sensors")

# 12. Agent = model + harness
s = slide("Dark: Empty", "12")
kicker(s, "02 · Implement", dark=True)
statement(s, [[("Agent = model + ", {}), ("harness", {"color": CYAN})]], size=72, color=WHITE, align=PP_ALIGN.CENTER)

# 13. Five runs
s = slide("Light: Empty", "13")
kicker(s, "Five agents with the same idea, prompt and model, and no human in the loop")
RUNS = (("raw-low", "no skills · low"), ("raw-med", "no skills · medium"), ("ecc-med", "ECC · medium"),
        ("firstmate-med", "firstmate · medium"), ("matt-med", "Pocock · medium"))
tw, th, gap = 2.25, 1.75, 0.15
x0, y = (W - (5 * tw + 4 * gap)) / 2, 2.0
mid = [x0 + i * (tw + gap) + tw / 2 for i in range(5)]
for i, (name, what) in enumerate(RUNS):
    x = x0 + i * (tw + gap)
    base = name == "raw-med"
    if base:
        text(s, x, y - 0.4, tw, 0.3, "BASELINE", size=12, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
    box(s, x, y, tw, th, fill=PAPER, line=CYAN if base else None, weight=2.5, name=f"Run {name}")
    text(s, x + 0.18, y + 0.25, tw - 0.36, 0.5, name, size=19, bold=True)
    text(s, x + 0.18, y + 0.8, tw - 0.36, 0.4, what, size=14, color=INK2)
    text(s, x + 0.18, y + 1.2, tw - 0.36, 0.4, f":300{i + 1}", size=14, color=MUTED)
# Brackets: what each comparison varies, both measured against the baseline
for (a, b), by, label in (((0, 1), 4.2, "effort"), ((1, 4), 5.15, "skill pack")):
    line(s, mid[a], by, mid[b], by, color=INK2, weight=1.5)
    for t in (a, b):
        line(s, mid[t], by - 0.18, mid[t], by, color=INK2, weight=1.5)
    text(s, mid[a], by + 0.1, mid[b] - mid[a], 0.45, label, size=20, bold=True, color=INK2, align=PP_ALIGN.CENTER)

# 14. What it costs: slopalytics' cost vs intelligence chart
s = slide("Light: Empty", "14")
text(s, M, 0.6, 12, 0.8, "Every step up in intelligence costs more per task", size=34, bold=True)
pic = s.shapes.add_picture(str(ASSETS / "slopalytics.png"), Inches(M), Inches(1.6), height=Inches(5.0))
pic.name = "slopalytics.com"
tx = M + pic.width / 914400 + 0.35
text(s, tx, 2.2, W - M - tx, 3.4, [[("Opus 5.5", {"bold": True, "color": INK, "size": 22})],
                                    [("$0.55", {"bold": True, "color": CYAN, "size": 40})],
                                    [("a task at low effort", {})],
                                    [("$5.98", {"bold": True, "color": CYAN, "size": 40})],
                                    [("at max", {})]], size=16, color=INK2)
text(s, M, 6.75, 12, 0.35, "slopalytics.com · data from Artificial Analysis", size=13, color=MUTED)

# 15. Guides and sensors
s = slide("Light: Empty", "15")
kicker(s, "02 · Implement")
for i, (ic, head, when, items) in enumerate((("compass", "Guides", "before it acts", "AGENTS.md · skills · conventions"),
                                             ("gauge", "Sensors", "after it acts", "tests · types · lint · the browser"))):
    x = M + 0.4 + i * 6.2
    icon(s, ic, x, 1.7, 1.3)
    text(s, x, 3.25, 5.4, 0.9, head, size=54, bold=True)
    text(s, x, 4.25, 5.4, 0.6, when, size=26, color=CYAN, bold=True)
    text(s, x, 5.0, 5.4, 0.8, items, size=20, color=INK2)

# 16. State machines: the process as a hard rule
s = slide("Light: Empty", "16")
kicker(s, "02 · Implement")
text(s, M, 0.9, 12, 0.8, "Or make the process a hard rule instead of asking nicely", size=32, bold=True)
STATES = (("plan", ("read", "search")), ("edit", ("read", "write")), ("test", ("npm test",)), ("done", ()))
bw, gap, by = 2.3, 0.85, 2.6
for i, (name, tools) in enumerate(STATES):
    x = M + 0.2 + i * (bw + gap)
    last = i == len(STATES) - 1
    box(s, x, by, bw, 1.0, fill=INK if last else PAPER, line=None if last else LINE, radius=0.5, name=f"State {name}")
    text(s, x, by, bw, 1.0, name, size=24, bold=True, color=WHITE if last else INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if not last:
        arrow(s, x + bw + 0.25, by + 0.37, 0.35, 0.26, color=MUTED)
    for j, t in enumerate(tools):
        box(s, x + 0.3, by + 1.2 + j * 0.45, bw - 0.6, 0.36, line=CYAN, radius=0.5, weight=1)
        text(s, x + 0.3, by + 1.2 + j * 0.45, bw - 0.6, 0.36, [[(t, {"font": MONO})]], size=12, color=INK2,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
# Red tests send it back to edit, a fixed number of times
ex, tx = M + 0.2 + (bw + gap) + bw / 2, M + 0.2 + 2 * (bw + gap) + bw / 2
line(s, tx, by, tx, by - 0.35, color=MUTED, dash=True)
line(s, tx, by - 0.35, ex, by - 0.35, color=MUTED, dash=True)
line(s, ex, by - 0.35, ex, by - 0.02, color=MUTED, dash=True, arrow=True)
text(s, ex, by - 0.72, tx - ex, 0.3, "red: back to edit, 3 tries", size=13, color=MUTED, align=PP_ALIGN.CENTER)
text(s, M, 5.35, 6.0, 1.0, [[("Statewright · Stately", {"bold": True, "color": INK})],
                            [("Enforced per state, not asked for in the prompt", {})]], size=18, color=INK2)
text(s, 6.6, 5.2, 6.13, 0.7, "local model: fail → 10/10", size=32, bold=True, color=CYAN, align=PP_ALIGN.RIGHT)
text(s, 6.6, 5.95, 6.13, 0.4, "13.8 GB model, 5 SWE-bench tasks, vendor's figures", size=14, color=MUTED, align=PP_ALIGN.RIGHT)

# 17. 40 vs 300
s = slide("Dark: Empty", "17")
kicker(s, "02 · Implement", dark=True)
for i, (num, who) in enumerate((("~40", "skills · Matt Pocock"), ("~300", "skills · ECC"))):
    x = M + i * 6.2
    text(s, x, 1.5, 6, 2.4, num, size=150, bold=True, color=WHITE if i == 0 else CYAN, anchor=MSO_ANCHOR.BOTTOM)
    text(s, x + 0.1, 4.0, 6, 0.6, who, size=26, color=GREY_ON_DARK)
text(s, M + 0.1, 5.6, 12, 0.6, [[("firstmate:  ", {"bold": True, "color": WHITE}), ("1 first mate → a crew", {})]], size=26, color=GREY_ON_DARK)

# 18. How each run is driven
s = slide("Light: Empty", "18")
kicker(s, "02 · Implement")
text(s, M, 0.9, 12, 0.9, "Each pack runs the way its author meant it to", size=36, bold=True)
cx = (M, 2.55, 6.35)
cwid = (1.8, 3.6, W - M - 6.35)
for x, w, h in zip(cx, cwid, ("Run", "What is installed", "What the agent is told, turn by turn")):
    text(s, x, 1.95, w, 0.35, h.upper(), size=12, bold=True, color=MUTED)
CMD, TO = {"font": MONO, "bold": True, "color": INK}, {"color": CYAN, "bold": True}
DRIVE = (("raw", "Nothing but AGENTS.md", [("The idea, once", {})]),
         ("ecc", "ECC core profile: skills, agents, hooks",
          [("/plan", CMD), ("  →  ", TO), ("CONFIRM", CMD), (" + tdd-workflow", {}), ("  →  ", TO), ("/code-review", CMD)]),
         ("firstmate", "A firstmate clone, the first mate",
          [("One request: add the game as a project and ship it; the crew does the rest", {})]),
         ("matt", "Pocock’s skills, setup done in advance",
          [("/grill-with-docs", CMD), ("  →  ", TO), ("/to-spec", CMD), ("  →  ", TO), ("/implement", CMD)]))
rh, ry = 0.9, 2.35
for i, (run, installed, told) in enumerate(DRIVE):
    y = ry + i * rh
    line(s, M, y, W - M, y, color=LINE, weight=1)
    text(s, cx[0], y, cwid[0], rh, run, size=22, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, cx[1], y, cwid[1], rh, installed, size=15, color=INK2, anchor=MSO_ANCHOR.MIDDLE)
    text(s, cx[2], y, cwid[2], rh, [told], size=16, color=INK2, anchor=MSO_ANCHOR.MIDDLE)
line(s, M, ry + 4 * rh, W - M, ry + 4 * rh, color=LINE, weight=1)
po = box(s, M, 6.2, 0.6, 0.6, fill=CYAN, shape=MSO_SHAPE.OVAL, name="Product owner")
text(s, M, 6.2, 0.6, 0.6, "PO", size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, M + 0.85, 6.2, 11, 0.6, "Every question for a human goes to a product-owner agent", size=20, color=INK2,
     anchor=MSO_ANCHOR.MIDDLE)

# 19. Fold the output
s = slide("Light: Empty", "19")
kicker(s, "02 · Implement")
text(s, M, 0.9, 12, 1.2, "Most of what an agent reads is tool output it only needed one line of", size=32, bold=True)
LOG = ["$ npm run test:e2e", "Running 24 tests using 2 workers", "✓ players › joins with a name (412ms)",
       "✓ players › reconnects within 10 s (2.1s)", "✓ rounds › countdown starts (2.6s)", "✓ fighting › attack hits 25 (1.2s)",
       "✓ bot › joins a lone player (3.0s)", "✗ zone › shrinks every interval", "  expected 130, received 120",
       "✓ phone › buttons move (0.9s)", "… 412 more lines"]
lx, ly, lw, lh = M, 2.45, 3.9, 2.9
box(s, lx, ly, lw, lh, fill=CARD_DARK, radius=0.04, name="Tool output")
text(s, lx + 0.2, ly + 0.15, lw - 0.4, lh - 0.3,
     [[(t, {"font": MONO, "color": RED if t.startswith(("✗", "  exp")) else GREY_ON_DARK})] for t in LOG], size=11)
f = box(s, 4.85, 3.15, 1.5, 1.5, fill=LINE, shape=MSO_SHAPE.FUNNEL, name="Fold")
f.rotation = 270.0
arrow(s, 6.45, 3.77, 0.3, 0.26)
box(s, 6.9, 3.6, 3.5, 0.6, fill=PAPER, radius=0.15, name="One line")
text(s, 6.9, 3.6, 3.5, 0.6, [[("✗ zone: expected 130, got 120", {"font": MONO})]], size=13, color=RED,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
arrow(s, 10.55, 3.77, 0.3, 0.26)
box(s, 11.0, 3.35, 1.73, 1.1, fill=INK, name="Model")
text(s, 11.0, 3.35, 1.73, 1.1, "model", size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
# The full log is kept, not thrown away
line(s, 5.6, 4.75, 5.6, 5.4, color=MUTED, dash=True, arrow=True)
box(s, 5.05, 5.45, 1.1, 0.95, fill=PAPER, line=LINE, shape=MSO_SHAPE.CAN, name="Archive")
text(s, 6.35, 5.5, 2.6, 0.9, [[("full log archived", {"bold": True, "color": INK})], [("read back by ID", {})]],
     size=15, color=INK2, anchor=MSO_ANCHOR.MIDDLE)
text(s, 9.0, 5.25, 3.73, 0.8, "331k → 76k", size=44, bold=True, color=CYAN, align=PP_ALIGN.RIGHT)
text(s, 8.0, 6.05, 4.73, 0.4, "tokens per call, 30-step session, vendor’s figures", size=13, color=MUTED, align=PP_ALIGN.RIGHT)
text(s, M, 6.6, 6, 0.4, "context-mode.com", size=18, bold=True, color=INK2)

# 20. Search, not read
s = slide("Light: Empty", "20")
kicker(s, "02 · Implement")
text(s, M, 0.9, 12, 1.2, "An agent that can search the code understands it with far fewer tokens", size=32, bold=True)


def meter(s, x, y, h, frac, color):
    box(s, x, y, 0.4, h, fill=None, line=LINE, radius=0.2, name="Token meter")
    box(s, x, y + h * (1 - frac), 0.4, h * frac, fill=color, radius=0.2, name="Tokens used")
    text(s, x - 0.3, y + h + 0.08, 1.0, 0.3, "tokens", size=12, color=MUTED, align=PP_ALIGN.CENTER)


# Reading: file after file, every line paid for
text(s, M, 2.2, 5, 0.5, "Reading", size=24, bold=True)
for i in range(10):
    fx, fy = M + (i % 5) * 0.85, 2.95 + (i // 5) * 1.05
    box(s, fx, fy, 0.62, 0.8, fill=PAPER, line=LINE, shape=MSO_SHAPE.FOLDED_CORNER, name=f"File {i + 1}")
text(s, M, 5.1, 4.5, 0.4, "open, read, open the next one…", size=15, color=INK2)
meter(s, 5.3, 2.75, 2.5, 0.92, RED)
line(s, W / 2 + 0.1, 2.3, W / 2 + 0.1, 5.4, color=LINE)
# Searching: one query into an index, a few exact hits back
sx = 7.25
text(s, sx, 2.2, 5, 0.5, "Searching", size=24, bold=True)
box(s, sx, 2.85, 3.9, 0.5, fill=PAPER, radius=0.25, name="Query")
text(s, sx, 2.85, 3.9, 0.5, [[("who calls joinRoom?", {"font": MONO})]], size=13, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
line(s, sx + 1.95, 3.37, sx + 1.95, 3.68, color=MUTED, arrow=True)
box(s, sx + 0.95, 3.7, 2.0, 0.85, fill=PAPER, line=LINE, shape=MSO_SHAPE.CAN, name="Index")
text(s, sx + 0.95, 3.85, 2.0, 0.65, "index of every repo", size=12, color=INK2, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
line(s, sx + 1.95, 4.57, sx + 1.95, 4.8, color=MUTED, arrow=True)
text(s, sx, 4.82, 3.9, 0.7, [[(h, {"font": MONO})] for h in ("game/room.ts:42", "lobby/api.ts:17", "web/join.tsx:8")],
     size=11, color=CYAN, align=PP_ALIGN.CENTER)
meter(s, 11.9, 2.75, 2.5, 0.18, CYAN)
# The two tools
for i, (name, what) in enumerate((("Sourcebot", "self-hosted code search, with an MCP server for your agent"),
                                  ("Sourcegraph", "indexes every repo; MCP server with precise code intelligence"))):
    x = M + i * 6.2
    text(s, x, 5.85, 5.9, 1.0, [[(name, {"size": 22, "bold": True, "color": INK})], [(what, {})]], size=15, color=INK2)

# 21. Section: QA
section("21", "03", "QA", "Prove it works, with numbers rather than vibes")

# 22. Our job has changed
s = slide("Dark: Basic Blue 2", "22")
kicker(s, "03 · QA", dark=True)
statement(s, [[("Code has become cheap, and our job is now ", {}), ("verifying its quality", {"color": CYAN_LIGHT})]],
          size=52, color=WHITE, w=8.4)

# 23. 8%
s = slide("Light: Empty", "23")
kicker(s, "03 · QA")
text(s, M, 1.3, 8, 3.8, "8%", size=240, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
text(s, M + 0.15, 5.25, 9, 0.9, [[("wrote tests first, before AI", {"size": 26, "color": INK2})],
                                 [("Diffblue developer survey, 2020", {"size": 14, "color": MUTED})]])

# 24. Red, green
s = slide("Light: Blob 4", "24")
kicker(s, "03 · QA")
for i, c in enumerate((RED, CYAN)):
    box(s, M + i * 0.75, 2.05, 0.5, 0.5, fill=c, shape=MSO_SHAPE.OVAL)
statement(s, "“Use red/green TDD.”", size=60, w=8.6, y=2.55)
text(s, M, 5.0, 8, 0.5, "Simon Willison", size=20, color=INK2)

# 25. How do we know?
s = slide("Light: Empty", "25")
text(s, M, 0.55, 12, 0.9, "How do we know which one is better?", size=40, bold=True)
layers = (("check", "Does it work?", "its own Playwright tests, then the hidden suite"),
          ("code", "Is the code any good?", "same tools and config for all, plus mutation testing"),
          ("coins", "What did it cost?", "tokens, dollars and minutes"),
          ("bot", "What does an AI reviewer think?", "blinded, a different model, and last"))
for i, (ic, head, sub) in enumerate(layers):
    y = 1.75 + i * 1.3
    last = i == 3
    box(s, M, y, W - 2 * M, 1.05, fill=PAPER if not last else None, line=LINE if last else None, name=f"Layer {i + 1}")
    icon(s, ic, M + 0.3, y + 0.22, 0.62)
    text(s, M + 1.3, y, 6.0, 1.05, head, size=24, bold=True, color=INK if not last else INK2, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M + 7.4, y, 4.5, 1.05, sub, size=18, color=INK2, anchor=MSO_ANCHOR.MIDDLE)

SCORE_ROWS = ["Acceptance (hidden suite)", "Mutation score", "Cost", "Minutes"]
SCORE_COLS = [name for name, _ in RUNS]

# 26. Effort results
s = slide("Light: Empty", "26")
kicker(s, "Result 1 · thinking effort")
text(s, M, 0.9, 12, 0.9, "What did more thinking buy us?", size=40, bold=True)
scorecard(s, SCORE_COLS, SCORE_ROWS, M, 2.4, W - 2 * M, first_w=3.4, head_size=15, label_size=20,
          dim=("ecc-med", "firstmate-med", "matt-med"), baseline="raw-med")

# 27. Skills results
s = slide("Light: Empty", "27")
kicker(s, "Result 2 · skills")
text(s, M, 0.9, 12, 0.9, "Which skill pack did best?", size=40, bold=True)
scorecard(s, SCORE_COLS, SCORE_ROWS, M, 2.4, W - 2 * M, first_w=3.4, head_size=15, label_size=20,
          dim=("raw-low",), baseline="raw-med")

# 28. Section: Deploy
section("28", "04", "Deploy", "Put it in front of real people, and watch the bill")

# 29. Play it
s = slide("Light: Empty", "29")
kicker(s, "04 · Deploy")
text(s, M, 0.55, 12, 0.9, "Grab your phone and play one of them", size=40, bold=True)
q, gap = 1.8, 0.55
x0 = (W - (5 * q + 4 * gap)) / 2
for i, n in enumerate(SCORE_COLS):
    x = x0 + i * (q + gap)
    box(s, x, 2.9, q, q, fill=PAPER, line=LINE, radius=0.04, name=f"QR {n}")
    text(s, x, 2.9, q, q, "QR", size=16, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x - 0.2, 4.7, q + 0.4, 0.5, n, size=20, bold=True, align=PP_ALIGN.CENTER)
    text(s, x - 0.2, 5.15, q + 0.4, 0.4, f":300{i + 1}", size=14, color=MUTED, align=PP_ALIGN.CENTER)

# 30. Business fit
s = slide("Dark: Picture - Right 3", "30")
statement(s, "When code is cheap, what matters is whether it fits the business and the people using it",
          size=40, color=WHITE, w=6.4)

# 31. At work
s = slide("Light: Empty", "31")
text(s, M, 0.55, 12, 0.9, "What happens at work", size=40, bold=True)
steps = (("rocket", "You ship faster", INK), ("coins", "The bill grows", INK), ("scissors", "So you cut context", CYAN))
cw, gap = 3.4, 0.9
x0 = (W - (3 * cw + 2 * gap)) / 2
for i, (ic, label, col) in enumerate(steps):
    x = x0 + i * (cw + gap)
    icon(s, ic, x + (cw - 1.2) / 2, 2.3, 1.2)
    text(s, x, 3.85, cw, 0.8, label, size=30, bold=True, color=col, align=PP_ALIGN.CENTER)
    if i < 2:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + cw + 0.2), Inches(2.7), Inches(0.5), Inches(0.4))
        a.fill.solid()
        a.fill.fore_color.rgb = LINE
        a.line.fill.background()

# 32. Takeaways: one per step
s = slide("Dark: Empty", "32")
kicker(s, "Takeaways", dark=True)
loop(s, 3.0, 3.85, 1.45, lit="all", dark=True, labels=("", "", "", ""))
TAKE = ("Spec what you want, not how to build it", "Let the agent implement inside a harness",
        "Prove it with QA you can measure", "Deploy it, and watch what it costs")
text(s, 5.2, 1.5, 7.6, 4.8, [[(f"{i}  ", {"color": CYAN}), (t, {})] for i, t in enumerate(TAKE, 1)],
     size=25, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, spacing=22)

# Appendix
# A1. Questions
s = slide("Dark: Title and Subtitle", "A1", keep_title=True)
title(s, "Questions?", size=60)
text(s, 2.0, 4.55, W - 4.0, 0.5, "Sakari Nahi · Zure", size=22, color=GREY_ON_DARK, align=PP_ALIGN.CENTER)

# A2. We've lost something: the photo carries it
s = slide("Dark: Picture - Left 2", "A2", keep_title=True)
title(s, "We’ve lost something: the flow", size=48)

# A3. What have we gained?
s = slide("Light: Blob 2", "A3")
statement(s, "What have we gained?", size=64, w=8.0)

# A4. Two Februaries
s = slide("Light: Empty", "A4")
y = 3.75
s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(M + 0.2), Inches(y), Inches(W - M - 0.2), Inches(y)).line.color.rgb = LINE
for x, year, word, on in ((2.2, "Feb 2025", "Vibe coding", False), (8.0, "Feb 2026", "Agentic engineering", True)):
    box(s, x, y - 0.22, 0.44, 0.44, fill=CYAN if on else WHITE, line=CYAN if on else MUTED, shape=MSO_SHAPE.OVAL, weight=2)
    text(s, x, y - 1.25, 5, 0.6, year, size=24, color=INK2)
    text(s, x, y + 0.5, 5, 1.0, word, size=40, bold=True, color=INK if on else INK2)

# A5. Determinism
s = slide("Dark: Blob 1", "A5")
statement(s, "What we’re really after is determinism", size=60, color=WHITE, w=8.0)

# A6. −19%
s = slide("Light: Empty", "A6")
kicker(s, "Appendix")
text(s, M, 1.3, 9, 3.8, "−19%", size=220, bold=True, color=RED, anchor=MSO_ANCHOR.MIDDLE)
text(s, M + 0.15, 5.25, 11, 0.9, [[("experienced developers with AI, while they believed +20%", {"size": 26, "color": INK2})],
                                 [("METR randomized trial, 2025", {"size": 14, "color": MUTED})]])

# A7. Waterfall?
s = slide("Light: Basic Red 2", "A7")
kicker(s, "Appendix")
statement(s, "Isn’t this just waterfall with markdown?", size=56, w=8.2)

# A8. Harnesses age
s = slide("Dark: Blob 3", "A8")
text(s, 5.2, 0.55, 6, 0.3, "APPENDIX", size=14, bold=True, color=CYAN_LIGHT)
statement(s, "Every harness encodes something the model can’t do yet", size=48, color=WHITE, x=5.2, w=7.4)

prs.save(OUT)
print(f"wrote {OUT} ({len(prs.slides)} slides)")
